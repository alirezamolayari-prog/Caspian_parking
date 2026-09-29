"""Import existing subscribers from Excel (SPEC §4.6) with a validation report.

Rows with any problem are skipped and listed in the report; all valid rows are imported together.
Imported start dates become a zero-amount ``import`` subscription payment (so revenue reports are
not inflated) and the row's amount becomes the person's subscription price.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Any

from sqlalchemy import select

from caspian_parking.core.digits import normalize_input, parse_int
from caspian_parking.core.jalali import local_to_utc, parse_jdate
from caspian_parking.core.permissions import Permission
from caspian_parking.core.plate import Plate, PlateKind, parse_plate
from caspian_parking.data.models import Person, PersonPlate, Shop, SubscriptionPayment
from caspian_parking.i18n import tr
from caspian_parking.services.context import AppContext
from caspian_parking.services.exporters import ExportTable, export_excel
from caspian_parking.services.people import PeopleError, normalize_mobile
from caspian_parking.services.tariff_service import load_schedule

FIELDS = (
    "first_name",
    "last_name",
    "mobile",
    "shop",
    "location",
    "vehicle_model",
    "plate1",
    "plate2",
    "start_date",
    "amount",
    "notes",
)
REQUIRED = frozenset({"first_name", "plate1"})


@dataclass(frozen=True)
class ImportIssue:
    row: int
    field: str
    message: str


@dataclass
class ImportReport:
    total_rows: int = 0
    imported: int = 0
    issues: list[ImportIssue] = field(default_factory=list)
    report_path: Path | None = None

    @property
    def skipped_rows(self) -> int:
        return len({issue.row for issue in self.issues})


@dataclass
class _Row:
    number: int
    first_name: str
    last_name: str
    mobile: str | None
    shop: str
    location: str
    vehicle_model: str | None
    plates: list[Plate]
    start: datetime | None
    amount: int | None
    notes: str | None


def write_template(path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = tr("import.sheet")
    sheet.sheet_view.rightToLeft = True
    for index, name in enumerate(FIELDS, start=1):
        cell = sheet.cell(row=1, column=index, value=tr(f"import.col_{name}"))
        cell.font = Font(name="Tahoma", bold=True)
        sheet.column_dimensions[cell.column_letter].width = 18
    example = [tr("import.example_first"), tr("import.example_last"), "09121234567", tr("import.example_shop"),
               tr("import.inside"), tr("import.example_model"), "12ب345-22", "", "1405/07/01", 4000000, ""]  # fmt: skip
    for index, value in enumerate(example, start=1):
        sheet.cell(row=2, column=index, value=value)
    help_sheet = workbook.create_sheet(tr("import.help_sheet"))
    help_sheet.sheet_view.rightToLeft = True
    for row, key in enumerate(("import.help_1", "import.help_2", "import.help_3", "import.help_4"), start=1):
        help_sheet.cell(row=row, column=1, value=tr(key))
    help_sheet.column_dimensions["A"].width = 110
    workbook.save(path)
    return path


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return normalize_input(str(value))


def _parse_row(number: int, values: dict[str, Any], issues: list[ImportIssue]) -> _Row | None:
    before = len(issues)

    def issue(field_name: str, key: str) -> None:
        issues.append(ImportIssue(number, field_name, tr(key)))

    texts = {name: _text(values.get(name)) for name in FIELDS}
    for name in REQUIRED:
        if not texts[name]:
            issue(name, "import.required")
    mobile = None
    if texts["mobile"]:
        try:
            mobile = normalize_mobile(texts["mobile"])
        except PeopleError:
            issue("mobile", "people.bad_mobile")
    plates: list[Plate] = []
    for name in ("plate1", "plate2"):
        if texts[name]:
            try:
                plates.append(parse_plate(texts[name], allow_free=False))
            except ValueError:
                issue(name, "import.bad_plate")
    start = None
    if texts["start_date"]:
        try:
            day = parse_jdate(texts["start_date"]).to_gregorian()
            start = local_to_utc(datetime.combine(day, time(0, 0)))
        except ValueError:
            issue("start_date", "import.bad_date")
    amount = None
    if texts["amount"]:
        try:
            amount = parse_int(texts["amount"])
            if amount < 0:
                raise ValueError
        except ValueError:
            issue("amount", "import.bad_amount")
    location = "outside" if texts["location"] in (tr("import.outside"), "outside", "خارج") else "inside"
    if len(issues) > before:
        return None
    return _Row(
        number,
        texts["first_name"],
        texts["last_name"],
        mobile,
        texts["shop"],
        location,
        texts["vehicle_model"] or None,
        plates,
        start,
        amount,
        texts["notes"] or None,
    )


def read_rows(path: Path) -> tuple[list[_Row], list[ImportIssue], int]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    issues: list[ImportIssue] = []
    rows: list[_Row] = []
    total = 0
    for number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(c not in (None, "") for c in cells):
            continue
        total += 1
        values = {name: cells[index] if index < len(cells) else None for index, name in enumerate(FIELDS)}
        parsed = _parse_row(number, values, issues)
        if parsed is not None:
            rows.append(parsed)
    workbook.close()
    return rows, issues, total


def import_subscribers(ctx: AppContext, path: Path, dry_run: bool = False) -> ImportReport:
    if not ctx.can(Permission.MANAGE_SUBSCRIBERS):
        raise PeopleError("gate.permission_denied")
    rows, issues, total = read_rows(path)
    report = ImportReport(total_rows=total, issues=issues)
    seen: dict[str, int] = {}
    valid: list[_Row] = []
    with ctx.read() as session:
        taken = set(session.scalars(select(PersonPlate.plate_key).where(PersonPlate.is_active.is_(True))))
    for row in rows:
        problems = False
        for plate in row.plates:
            if plate.key in taken:
                issues.append(ImportIssue(row.number, "plate", tr("people.plate_taken")))
                problems = True
            elif plate.key in seen:
                issues.append(ImportIssue(row.number, "plate", tr("import.duplicate_in_file", row=seen[plate.key])))
                problems = True
            else:
                seen[plate.key] = row.number
        if not problems:
            valid.append(row)
    if not dry_run and valid:
        with ctx.uow(reason=tr("import.reason")) as session:
            shops = {s.name: s.id for s in session.scalars(select(Shop).where(Shop.is_active.is_(True)))}
            days = load_schedule(session).at(ctx.clock.now_utc()).subscription_days
            for row in valid:
                person = Person(
                    kind="subscriber",
                    first_name=row.first_name,
                    last_name=row.last_name,
                    mobile=row.mobile,
                    shop_id=shops.get(row.shop),
                    brand=row.shop or None,
                    location_type=row.location,
                    vehicle_model=row.vehicle_model,
                    notes=row.notes,
                    price_override=row.amount,
                    negative_max_days=10,
                )
                session.add(person)
                session.flush()
                for plate in row.plates:
                    description = tr("vehicle.motorcycle") if plate.kind is PlateKind.MOTORCYCLE else None
                    session.add(PersonPlate(person_id=person.id, plate_key=plate.key, description=description))
                if row.start is not None:
                    end = row.start + timedelta(days=days)
                    session.add(
                        SubscriptionPayment(
                            person_id=person.id,
                            amount=0,
                            method="import",
                            paid_at_utc=row.start,
                            period_days=days,
                            new_end_utc=end,
                            used_days=[],
                        )
                    )
                    person.subscription_end_utc = end
        report.imported = len(valid)
    report.issues.sort(key=lambda i: (i.row, i.field))
    if report.issues:
        table = ExportTable.simple(
            tr("import.report_title"),
            [tr("import.col_row"), tr("import.col_field"), tr("import.col_problem")],
            [[i.row, tr(f"import.col_{i.field}") if i.field in FIELDS else i.field, i.message] for i in report.issues],
            widths=[10, 20, 60],
        )
        stamp = ctx.clock.now_utc().strftime("%Y%m%d-%H%M%S")
        report.report_path = export_excel(table, ctx.data_root.exports / f"import-report-{stamp}.xlsx")
    return report

"""Right-to-left, print-ready table exports: Excel (openpyxl) and Word (python-docx).

Used by the follow-up list and by every report (Phase 5). Plates are exported as text inside a
left-to-right embedding so their digit order never flips in Excel/Word.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from caspian_parking.core.digits import to_persian_digits
from caspian_parking.core.plate import Plate, PlateKind

EXPORT_FONT = "Tahoma"  # installed on every Windows PC and shapes Persian well
LRE = chr(0x202A)
PDF = chr(0x202C)
HEADER_FILL = "DCE9E8"
SECTION_FILL = "F2F2F2"


@dataclass
class Section:
    title: str
    rows: list[list[Any]]


@dataclass
class ExportTable:
    title: str
    columns: list[str]
    sections: list[Section] = field(default_factory=list)
    subtitle: str = ""
    widths: list[int] = field(default_factory=list)
    footer: str = ""

    @classmethod
    def simple(cls, title: str, columns: list[str], rows: list[list[Any]], **kwargs: Any) -> ExportTable:
        return cls(title, columns, [Section("", rows)], **kwargs)

    def row_count(self) -> int:
        return sum(len(section.rows) for section in self.sections)


def plate_text(plate: Plate | None) -> str:
    """``۱۲ ب ۳۴۵ - ۲۲`` in a left-to-right embedding (correct order inside RTL documents)."""
    if plate is None:
        return ""
    if plate.kind is PlateKind.CAR:
        left, letter, mid, region = plate.parts
        text = f"{left} {letter} {mid} - {region}"
    elif plate.kind is PlateKind.MOTORCYCLE:
        text = f"{plate.parts[0]}-{plate.parts[1]}"
    else:
        text = plate.parts[0]
    return f"{LRE}{to_persian_digits(text)}{PDF}"


def export_excel(table: ExportTable, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = table.title[:31] or "Sheet"
    sheet.sheet_view.rightToLeft = True
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base = Font(name=EXPORT_FONT, size=11)
    bold = Font(name=EXPORT_FONT, size=11, bold=True)
    title_font = Font(name=EXPORT_FONT, size=14, bold=True)
    right = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    columns = len(table.columns)
    sheet.cell(row=1, column=1, value=table.title).font = title_font
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, columns))
    row = 2
    if table.subtitle:
        sheet.cell(row=row, column=1, value=table.subtitle).font = base
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(1, columns))
        row += 1
    header_row = row + 1
    for index, title in enumerate(table.columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = bold
        cell.alignment = center
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    row = header_row + 1
    for section in table.sections:
        if section.title:
            cell = sheet.cell(row=row, column=1, value=section.title)
            cell.font = bold
            cell.fill = PatternFill("solid", fgColor=SECTION_FILL)
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(1, columns))
            row += 1
        for values in section.rows:
            for index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=index, value=value)
                cell.font = base
                cell.alignment = right
                cell.border = border
                if isinstance(value, int):
                    cell.number_format = "#,##0"
            sheet.row_dimensions[row].height = 26
            row += 1
    if table.footer:
        sheet.cell(row=row + 1, column=1, value=table.footer).font = base
    for index in range(1, columns + 1):
        width = table.widths[index - 1] if index - 1 < len(table.widths) else 18
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)
    return path


def _rtl_paragraph(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    properties.append(bidi)


def _rtl_run(run: Any, bold: bool = False, size: int = 10) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = EXPORT_FONT
    run.font.size = Pt(size)
    run.bold = bold
    properties = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    properties.append(rtl)
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:cs"), EXPORT_FONT)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(size * 2))
    properties.append(size_cs)
    if bold:
        properties.append(OxmlElement("w:bCs"))


def _cell_text(cell: Any, text: str, bold: bool = False, size: int = 10) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _rtl_paragraph(paragraph)
    _rtl_run(paragraph.add_run(text), bold=bold, size=size)


def export_word(table: ExportTable, path: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Cm(1.5))
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rtl_paragraph(heading)
    _rtl_run(heading.add_run(table.title), bold=True, size=16)
    if table.subtitle:
        sub = document.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _rtl_paragraph(sub)
        _rtl_run(sub.add_run(table.subtitle), size=11)
    grid = document.add_table(rows=1, cols=len(table.columns))
    grid.style = "Table Grid"
    table_properties = grid._tbl.tblPr
    visual = OxmlElement("w:bidiVisual")
    visual.set(qn("w:val"), "1")
    table_properties.append(visual)
    header = grid.rows[0]
    for index, title in enumerate(table.columns):
        _cell_text(header.cells[index], title, bold=True)
    header_properties = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for part in table.sections:
        if part.title:
            cells = grid.add_row().cells
            merged = cells[0].merge(cells[-1])
            _cell_text(merged, part.title, bold=True, size=11)
        for values in part.rows:
            cells = grid.add_row().cells
            for index, value in enumerate(values):
                text = (
                    ""
                    if value is None
                    else (
                        to_persian_digits(f"{value:,}".replace(",", "٬"))
                        if isinstance(value, int) and not isinstance(value, bool)
                        else str(value)
                    )
                )
                _cell_text(cells[index], text)
    if table.footer:
        footer = document.add_paragraph()
        _rtl_paragraph(footer)
        _rtl_run(footer.add_run(table.footer), size=9)
    document.save(path)
    return path

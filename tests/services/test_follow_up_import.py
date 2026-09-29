"""Follow-up list, Excel/Word exports and Excel import of subscribers."""

from __future__ import annotations

import shutil
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from caspian_parking.core.jalali import local_to_utc
from caspian_parking.core.plate import parse_plate
from caspian_parking.core.subscriptions import Light
from caspian_parking.data.models import Person, PersonPlate, SubscriptionPayment
from caspian_parking.services import auth
from caspian_parking.services.auth import CurrentUser
from caspian_parking.services.exporters import ExportTable, Section, export_excel, export_word, plate_text
from caspian_parking.services.follow_up import export_follow_up, follow_up_rows, follow_up_table
from caspian_parking.services.people import PeopleError, PeopleService, PersonInput
from caspian_parking.services.subscriber_import import FIELDS, import_subscribers, write_template

MON = date(2026, 9, 28)
ARTIFACTS = Path(__file__).resolve().parents[1] / "artifacts"


def at(day: date, hh: int = 10) -> datetime:
    return local_to_utc(datetime.combine(day, time(hh, 0)))


@pytest.fixture
def svc(app_ctx, clock):
    clock.set(at(MON))
    with app_ctx.uow() as session:
        user = auth.create_user(session, "boss", "Boss", "secret1", preset="admin")
    app_ctx.user = CurrentUser.from_user(user)
    return PeopleService(app_ctx)


def _person(svc, name, plate, days_until_end):
    person = svc.create_person(PersonInput(first_name=name, mobile="09121234567", plates=[(parse_plate(plate), None)]))
    svc.pay_subscription(person.id, "cash")
    with svc.ctx.uow() as session:
        stored = session.get(Person, person.id)
        stored.subscription_end_utc = at(MON) + timedelta(days=days_until_end)
    return person


def test_plate_text_keeps_ltr_order():
    text = plate_text(parse_plate("12ب345-22"))
    assert text.startswith(chr(0x202A))
    assert "۱۲ ب ۳۴۵ - ۲۲" in text
    assert plate_text(None) == ""
    assert "۱۲۳-۴۵۶۷۸" in plate_text(parse_plate("123-45678"))


def test_follow_up_sections_and_order(svc, app_ctx):
    _person(svc, "سبز", "11ب111-11", 20)
    _person(svc, "زرد", "12ب111-11", 4)
    _person(svc, "قرمز", "13ب111-11", 1)
    _person(svc, "سیاه‌کم", "14ب111-11", -2)
    _person(svc, "سیاه‌زیاد", "15ب111-11", -9)
    rows = follow_up_rows(app_ctx)
    assert [r.person.first_name for r in rows] == ["سیاه‌زیاد", "سیاه‌کم", "قرمز", "زرد"]
    assert [r.light for r in rows] == [Light.BLACK, Light.BLACK, Light.RED, Light.AMBER]
    table = follow_up_table(app_ctx, rows)
    assert [len(s.rows) for s in table.sections] == [2, 1, 1]
    assert len(table.columns) == 8
    assert table.sections[0].rows[0][7] == ""  # empty signature column


def test_follow_up_exports_open(svc, app_ctx, tmp_path):
    _person(svc, "علی", "12ب345-22", -3)
    xlsx = export_follow_up(app_ctx, tmp_path, "excel")
    docx = export_follow_up(app_ctx, tmp_path, "word")
    sheet = load_workbook(xlsx).active
    assert sheet.sheet_view.rightToLeft
    values = [c for row in sheet.iter_rows(values_only=True) for c in row if c]
    assert "علی" in values
    document = Document(str(docx))
    texts = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "علی" in texts
    ARTIFACTS.mkdir(exist_ok=True)
    shutil.copyfile(xlsx, ARTIFACTS / "follow_up.xlsx")
    shutil.copyfile(docx, ARTIFACTS / "follow_up.docx")


def test_generic_exporters(tmp_path):
    table = ExportTable(
        "گزارش", ["الف", "مبلغ"], [Section("بخش", [["x", 190000], ["y", None]])], subtitle="زیر", footer="پا"
    )
    excel = export_excel(table, tmp_path / "t.xlsx")
    word = export_word(table, tmp_path / "t.docx")
    cells = [c for row in load_workbook(excel).active.iter_rows(values_only=True) for c in row]
    assert 190000 in cells  # numbers stay numeric in Excel
    assert "۱۹۰٬۰۰۰" in [c.text for t in Document(str(word)).tables for r in t.rows for c in r.cells]
    assert ExportTable.simple("t", ["a"], [[1], [2]]).row_count() == 2


# ---------------------------------------------------------------- import


def _workbook(path: Path, rows: list[list]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(FIELDS))
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_template_is_readable(tmp_path):
    path = write_template(tmp_path / "template.xlsx")
    sheet = load_workbook(path).worksheets[0]
    assert sheet.max_column == len(FIELDS)
    assert sheet.cell(row=2, column=7).value == "12ب345-22"


def test_import_valid_and_invalid_rows(svc, app_ctx, tmp_path):
    svc.create_shop("مبلمان آرتا")
    rows = [
        [
            "علی",
            "رضایی",
            "09121234567",
            "مبلمان آرتا",
            "داخل",
            "پژو",
            "12ب345-22",
            "123-45678",
            "1405/07/01",
            3500000,
            "",
        ],
        ["", "بی‌نام", "", "", "", "", "11ب111-11", "", "", "", ""],  # missing first name
        ["سارا", "", "0912", "", "", "", "bad", "", "1405/13/01", "-5", ""],  # 4 problems
        ["مینا", "", "", "", "خارج", "", "12ب345-22", "", "", "", ""],  # duplicate plate in file
        ["رضا", "", "", "", "", "", "۲۲ ج ۵۵۵ ایران ۴۴", "", "", "", "یادداشت"],
    ]
    report = import_subscribers(app_ctx, _workbook(tmp_path / "in.xlsx", rows))
    assert report.total_rows == 5
    assert report.imported == 2
    assert report.skipped_rows == 3
    assert report.report_path is not None
    assert report.report_path.is_file()
    with app_ctx.read() as session:
        ali = session.scalar(select(Person).where(Person.first_name == "علی"))
        assert ali.shop_id is not None
        assert ali.price_override == 3_500_000
        assert ali.subscription_end_utc == local_to_utc(datetime(2026, 9, 23)) + timedelta(days=30)
        assert len(session.scalars(select(PersonPlate).where(PersonPlate.person_id == ali.id)).all()) == 2
        payment = session.scalars(select(SubscriptionPayment)).one()
        assert (payment.amount, payment.method) == (0, "import")
        reza = session.scalar(select(Person).where(Person.first_name == "رضا"))
        assert reza.notes == "یادداشت"
        assert reza.subscription_end_utc is None


def test_import_dry_run_and_existing_plates(svc, app_ctx, tmp_path):
    svc.create_person(PersonInput(first_name="قبلی", plates=[(parse_plate("12ب345-22"), None)]))
    path = _workbook(
        tmp_path / "in.xlsx",
        [
            ["علی", "", "", "", "", "", "12ب345-22", "", "", "", ""],
            ["رضا", "", "", "", "", "", "13ب345-22", "", "", "", ""],
        ],
    )
    dry = import_subscribers(app_ctx, path, dry_run=True)
    assert dry.imported == 0
    assert dry.skipped_rows == 1
    real = import_subscribers(app_ctx, path)
    assert real.imported == 1


def test_import_needs_permission(app_ctx, tmp_path):
    app_ctx.user = None
    with pytest.raises(PeopleError, match="permission"):
        import_subscribers(app_ctx, tmp_path / "x.xlsx")
